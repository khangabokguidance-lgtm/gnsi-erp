#!/usr/bin/env python3
"""
GNSI Premium Booklet Builder
Guidance Navodaya and Sainik Institute
Design: Sir Moirangthem Himan Singh, Founder

Turns a chapter question file (.docx) into a premium bilingual booklet:
cover page, contents, two-column questions sorted by type, answer key and
worked solutions. English in Calibri/Cambria, Manipuri in Bmei04.

USAGE
    pip install python-docx
    python gnsi_booklet.py  chapter.json
    python gnsi_booklet.py  chapter.json  --columns 1      (single column)

INPUT FILE FORMAT (.docx) - the same layout as your existing chapter files:
    Q1. English question                      <- bold, starts with "Q<number>."
    Manipuri line                             <- Bmei04 runs (numbers may be Arial)
    (A) ...  (B) ...  (C) ...  (D) ...        <- one or more lines
    ...
    Q1. B   Q2. C   Q3. D ...                 <- answer key (optional, at the end)

SETTINGS FILE (chapter.json) - see the example factors_and_multiples.json.
Only "source", "output" and "title" are required; everything else is optional.
"""
import sys, re, json, argparse
import docx
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import parse_xml, OxmlElement

# =============================================================================
#  DESIGN CODE  -  change the look of every booklet here
# =============================================================================
DESIGN = {
    # Brand colours (hex, no #)
    "navy":       "1B2A4A",   # primary: cover block, headings, section bars
    "gold":       "B8892B",   # accent: rules, option letters, type badge
    "gold_light": "D8C08A",   # small text on navy
    "cream":      "F3E6C4",   # Manipuri title on cover
    "ink":        "222222",   # body text
    "ink_mm":     "333333",   # Manipuri lines
    "muted":      "6B7280",   # captions, footer
    "tint":       "F4EFE3",   # alternate table rows
    "rule":       "D9D2C0",   # table borders, footer rule

    # Typefaces
    "font_head":  "Cambria",  # titles, headings, question numbers
    "font_body":  "Calibri",  # English text
    "font_mm":    "Bmei04",   # Manipuri (Meitei Mayek) text

    # Type sizes (pt)
    "size_title": 34, "size_body": 11,
    "size_q": 10, "size_mm": 11, "size_opt": 9.5,          # question pages
    "size_bar": 11.5,                                       # section bar title

    # Page (A4, cm)
    "page_w": 21.0, "page_h": 29.7,
    "margin_lr": 2.0, "margin_top": 2.2, "margin_bottom": 2.0,
    "column_gap": 0.8, "question_indent": 0.75,

    # Branding text
    "institute":  "Guidance Navodaya and Sainik Institute",
    "motto":      "Excellence  ·  Discipline  ·  Guidance",
    "author":     "Sir Moirangthem Himan Singh",
    "author_role":"Founder, Guidance Navodaya and Sainik Institute",
    "footer":     "Prepared by Sir Moirangthem Himan Singh, Founder",
    "audience":   "For Sainik School and Navodaya Vidyalaya entrance preparation",
}
D = DESIGN

# =============================================================================
#  1. READ THE SOURCE FILE
# =============================================================================
def _dec(t):
    """Bmei04 text is often stored as Symbol private-use characters (U+F0xx)."""
    return ''.join(chr(ord(c) - 0xF000) if 0xF000 <= ord(c) < 0xF100 else c for c in t)

def read_source(path):
    src = docx.Document(path); ps = src.paragraphs
    starts, key = {}, {}
    for i, p in enumerate(ps):
        m = re.match(r'\s*Q(\d+)\.\s', p.text)
        if m and p.runs and p.runs[0].bold and int(m.group(1)) not in starts:
            starts[int(m.group(1))] = i
        for n, a in re.findall(r'Q(\d+)\.\s*([A-D])\b(?!\w)', p.text):
            if not (p.runs and p.runs[0].bold): key[int(n)] = a
    nums = sorted(starts); Q = {}
    for k, n in enumerate(nums):
        i = starts[n]
        end = starts[nums[k + 1]] if k + 1 < len(nums) else len(ps)
        block = [p for p in ps[i:end] if p.text.strip()]
        eng = re.sub(r'^\s*Q\d+\.\s*', '', block[0].text).strip()
        mm, opt_lines = [], []
        for p in block[1:]:
            if re.match(r'\s*\(A\)', p.text) or (opt_lines and re.match(r'\s*\([B-D]\)', p.text)):
                opt_lines.append(p.text)
            elif opt_lines:
                break                                   # anything after options (e.g. key) ends the block
            elif not mm:
                mm = [(_dec(r.text), r.font.name == D["font_mm"]) for r in p.runs if r.text]
        opts = [o.strip() for o in re.split(r'\([A-D]\)', '\t'.join(opt_lines))[1:]][:4]
        Q[n] = dict(eng=eng, mm=mm, opts=opts)
    return Q, key

# =============================================================================
#  2. APPLY CORRECTIONS FROM THE SETTINGS FILE
# =============================================================================
def apply_edits(Q, cfg):
    for n, text in cfg.get("set_english", {}).items(): Q[int(n)]["eng"] = text
    for n, pairs in cfg.get("replace_english", {}).items():
        for old, new in pairs: Q[int(n)]["eng"] = Q[int(n)]["eng"].replace(old, new)
    for n, opts in cfg.get("set_options", {}).items(): Q[int(n)]["opts"] = list(opts)
    for n, ch in cfg.get("set_option", {}).items():
        for letter, text in ch.items(): Q[int(n)]["opts"]["ABCD".index(letter)] = text
    for n, edits in cfg.get("replace_manipuri", {}).items():
        runs = Q[int(n)]["mm"]
        for e in edits:                                # [find, replace, which_occurrence(1-based, optional)]
            old, new, nth = e[0], e[1], (e[2] if len(e) > 2 else 1)
            seen = 0
            for k, (t, b) in enumerate(runs):
                c = t.count(old)
                if seen + c >= nth:
                    pos = -1
                    for _ in range(nth - seen): pos = t.index(old, pos + 1)
                    runs[k] = (t[:pos] + new + t[pos + len(old):], b); break
                seen += c
    # tidy multiplication signs: "2x2" / "2 x 3" -> "2 × 3"
    for q in Q.values():
        q["opts"] = [re.sub(r'\s*x\s*', ' × ', o) if re.search(r'\d\s*x\s*\d', o) else o for o in q["opts"]]
        q["eng"] = re.sub(r'(\d) x (\d)', r'\1 × \2', q["eng"])

# =============================================================================
#  3. LOW-LEVEL WORD HELPERS
# =============================================================================
class W:
    def __init__(self, d): self.d = d
    @staticmethod
    def font(r, name, size=None, bold=None, color=None, italic=None, caps=False, spacing=None):
        r.font.name = name
        rpr = r._element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
        for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(k), name)
        if size: r.font.size = Pt(size)
        if bold is not None: r.bold = bold
        if italic is not None: r.italic = italic
        if color: r.font.color.rgb = RGBColor.from_string(color)
        if caps: r.font.all_caps = True
        if spacing is not None:
            s = OxmlElement('w:spacing'); s.set(qn('w:val'), str(spacing)); rpr.append(s)
        return r
    @staticmethod
    def shade(pr, fill):
        s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill); pr.append(s)
    @staticmethod
    def border(p, side, color, sz=12, space=4):
        pPr = p._p.get_or_add_pPr(); b = pPr.find(qn('w:pBdr'))
        if b is None: b = OxmlElement('w:pBdr'); pPr.append(b)
        e = OxmlElement(f'w:{side}')
        for k, v in (('val', 'single'), ('sz', sz), ('space', space), ('color', color)): e.set(qn(f'w:{k}'), str(v))
        b.append(e)
    def para(self, text='', font=None, size=11, bold=False, color=None, align=None, before=0, after=0,
             italic=False, caps=False, spacing=None, container=None):
        p = (container or self.d).add_paragraph()
        p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
        if align is not None: p.alignment = align
        if text: self.font(p.add_run(text), font or D["font_body"], size, bold, color or D["ink"], italic, caps, spacing)
        return p
    @staticmethod
    def page_field(p, size=9, color=None):
        r = p.add_run(); W.font(r, D["font_body"], size, color=color or D["navy"])
        for kind, txt in (('begin', 0), ('instr', 'PAGE'), ('separate', 0), ('text', '1'), ('end', 0)):
            if kind == 'instr':
                e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt
            elif kind == 'text':
                e = OxmlElement('w:t'); e.text = txt
            else:
                e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), kind)
            r._element.append(e)
    @staticmethod
    def cell_pad(cell, t=0, b=0, l=0, r=0):
        tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement('w:tcMar')
        for k, v in (('top', t), ('bottom', b), ('start', l), ('end', r)):
            e = OxmlElement(f'w:{k}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
        tcPr.append(m)
    @staticmethod
    def tbl_borders(tbl, color=None, none=False):
        ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        inner = ''.join((f'<w:{e} w:val="nil"/>' if none else f'<w:{e} w:val="single" w:sz="4" w:color="{color or D["rule"]}"/>')
                        for e in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV'])
        tbl._tbl.tblPr.append(parse_xml(f'<w:tblBorders {ns}>{inner}</w:tblBorders>'))
    @staticmethod
    def widths(tbl, ws):
        tbl.autofit = False
        for gc, w in zip(tbl._tbl.tblGrid.findall(qn('w:gridCol')), ws): gc.set(qn('w:w'), str(int(Emu(int(w)).twips)))
        for row in tbl.rows:
            for c, w in zip(row.cells, ws): c.width = Emu(int(w))
    def section(self, cols):
        s = self.d.add_section(WD_SECTION.NEW_PAGE); s.different_first_page_header_footer = False
        sp = s._sectPr
        for c in sp.findall(qn('w:cols')): sp.remove(c)
        c = OxmlElement('w:cols'); c.set(qn('w:num'), str(cols)); c.set(qn('w:space'), str(int(Cm(D["column_gap"]).twips)))
        if cols > 1: c.set(qn('w:sep'), '1')
        sp.append(c)
    def mm(self, p, runs, size):
        for t, is_mm in runs:
            self.font(p.add_run(t), D["font_mm"] if is_mm else D["font_body"], size if is_mm else size - 1, False, D["ink_mm"])
    def bullet(self, text):
        p = self.para(after=3); p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.first_line_indent = Cm(-0.4)
        self.font(p.add_run('■  '), D["font_body"], 7, color=D["gold"]); self.font(p.add_run(text), D["font_body"], 10.5, color=D["ink"])

# =============================================================================
#  4. BUILD THE BOOKLET
# =============================================================================
def build(cfg, columns=2):
    Q, key = read_source(cfg["source"])
    apply_edits(Q, cfg)
    ANS = {**key, **{int(k): v for k, v in cfg.get("answers", {}).items()}}
    SOL = {int(k): v for k, v in cfg.get("solutions", {}).items()}
    groups = cfg.get("groups") or [{"code": "A", "name": "Practice Questions", "questions": sorted(Q)}]
    missing = sorted(set(Q) - {n for g in groups for n in g["questions"]})
    if missing: groups.append({"code": chr(ord(groups[-1]["code"]) + 1), "name": "Other Questions", "questions": missing})

    d = docx.Document(); w = W(d)
    st = d.styles['Normal']; st.font.name = D["font_body"]; st.font.size = Pt(D["size_body"])
    st.font.color.rgb = RGBColor.from_string(D["ink"]); st.element.rPr.rFonts.set(qn('w:eastAsia'), D["font_body"])
    st.paragraph_format.space_after = Pt(0); st.paragraph_format.line_spacing = 1.1
    s = d.sections[0]
    s.page_width, s.page_height = Cm(D["page_w"]), Cm(D["page_h"])
    s.left_margin = s.right_margin = Cm(D["margin_lr"]); s.top_margin = Cm(D["margin_top"]); s.bottom_margin = Cm(D["margin_bottom"])
    s.header_distance = s.footer_distance = Cm(1.0)
    CW = s.page_width - s.left_margin - s.right_margin
    COLW = int((CW - Cm(D["column_gap"])) / 2) if columns == 2 else CW
    IND = Cm(D["question_indent"])
    title = cfg["title"]

    # header & footer (not on cover)
    s.different_first_page_header_footer = True
    hp = s.header.paragraphs[0]; hp.style = st; hp.paragraph_format.tab_stops.add_tab_stop(CW, WD_TAB_ALIGNMENT.RIGHT)
    w.font(hp.add_run(D["institute"]), D["font_head"], 9, True, D["navy"], caps=True, spacing=10)
    w.font(hp.add_run('\t' + title), D["font_head"], 9, False, D["gold"], italic=True)
    w.border(hp, 'bottom', D["gold"], 6, 4)
    fp = s.footer.paragraphs[0]; fp.style = st; fp.paragraph_format.tab_stops.add_tab_stop(CW, WD_TAB_ALIGNMENT.RIGHT)
    w.border(fp, 'top', D["rule"], 4, 6)
    w.font(fp.add_run(D["footer"]), D["font_body"], 8.5, color=D["muted"], italic=True)
    w.font(fp.add_run('\tPage '), D["font_body"], 9, color=D["muted"]); w.page_field(fp)

    # ---- cover ----
    C = WD_ALIGN_PARAGRAPH.CENTER
    w.para(before=40)
    w.para(D["institute"].upper(), D["font_head"], 12, True, D["gold"], C, spacing=40)
    w.para(D["motto"], D["font_body"], 9.5, False, D["muted"], C, before=4, italic=True)
    w.para(before=60)
    t = d.add_table(rows=1, cols=1); w.tbl_borders(t, none=True); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]; w.shade(c._tc.get_or_add_tcPr(), D["navy"]); w.cell_pad(c, 500, 500, 400, 400); w.widths(t, [CW])
    cp = c.paragraphs[0]; cp.alignment = C
    w.font(cp.add_run(cfg.get("label", "MATHEMATICS").upper()), D["font_body"], 10, True, D["gold_light"], spacing=30)
    w.para(title, D["font_head"], D["size_title"], True, "FFFFFF", C, before=10, container=c)
    if cfg.get("title_mm"):
        p = w.para(container=c, align=C, before=6); w.font(p.add_run(cfg["title_mm"]), D["font_mm"], 20, color=D["cream"])
    w.para(f'{len(Q)} objective questions  ·  English with Manipuri', D["font_body"], 12, False, "E5E7EB", C, before=16, container=c)
    w.para('Sorted by type  ·  Answer key' + ('  ·  Worked solutions' if SOL else ''), D["font_body"], 11, False, "C9CED8", C, before=2, container=c)
    w.para(before=50)
    p = w.para(align=C); w.border(p, 'top', D["gold"], 12, 1); p.paragraph_format.left_indent = p.paragraph_format.right_indent = Cm(6)
    w.para('PREPARED BY', D["font_body"], 9, True, D["muted"], C, before=10, spacing=40)
    w.para(D["author"], D["font_head"], 20, True, D["navy"], C, before=6)
    w.para(D["author_role"], D["font_body"], 11.5, False, D["gold"], C, before=4, italic=True)
    w.para(before=70)
    w.para(D["audience"], D["font_body"], 9.5, False, D["muted"], C, italic=True)
    d.add_page_break()

    # ---- contents ----
    w.para('Contents', D["font_head"], 22, True, D["navy"], after=2)
    p = w.para(after=14); w.border(p, 'bottom', D["gold"], 12, 2); p.paragraph_format.right_indent = Cm(13)
    tbl = d.add_table(rows=1, cols=4); w.tbl_borders(tbl); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cc, h in zip(tbl.rows[0].cells, ['Type', 'Topic', 'Questions', 'Count']):
        w.shade(cc._tc.get_or_add_tcPr(), D["navy"]); w.cell_pad(cc, 90, 90, 120, 120)
        w.font(cc.paragraphs[0].add_run(h), D["font_body"], 10, True, "FFFFFF")
    num, RANGE = 1, {}
    for k, g in enumerate(groups):
        RANGE[g["code"]] = (num, num + len(g["questions"]) - 1); num += len(g["questions"])
        vals = [g["code"], g["name"], f'{RANGE[g["code"]][0]} – {RANGE[g["code"]][1]}', str(len(g["questions"]))]
        for j, (cc, v) in enumerate(zip(tbl.add_row().cells, vals)):
            if k % 2: w.shade(cc._tc.get_or_add_tcPr(), D["tint"])
            w.cell_pad(cc, 80, 80, 120, 120)
            w.font(cc.paragraphs[0].add_run(v), D["font_head"] if j == 0 else D["font_body"], 11 if j == 0 else 10.5, j == 0, D["navy"] if j == 0 else D["ink"])
    for j, (cc, v) in enumerate(zip(tbl.add_row().cells, ['', 'Total', '', str(len(Q))])):
        w.cell_pad(cc, 80, 80, 120, 120); w.font(cc.paragraphs[0].add_run(v), D["font_body"], 10.5, True, D["navy"])
    w.widths(tbl, [Cm(1.6), Cm(9.4), Cm(3.4), Cm(2.6)])
    w.para(before=18); w.para('How to use this booklet', D["font_head"], 13, True, D["navy"], after=4)
    for line in cfg.get("instructions", [
        "Each question is given in English, with the Manipuri (Meitei Mayek) line below it in the Bmei04 font.",
        "Questions are grouped by type, so a student can practise one skill at a time.",
        "Try every question in a group before turning to the answer key; then read the worked solution for any you missed.",
        "The Manipuri lines need the Bmei04 font installed on the computer to display correctly."]):
        w.bullet(line)

    # ---- questions ----
    w.section(columns)
    NEW, num = {}, 1
    for g in groups:
        p = w.para(before=10, after=8); p.paragraph_format.keep_with_next = True
        w.shade(p._p.get_or_add_pPr(), D["navy"])
        for side in ('top', 'bottom'): w.border(p, side, D["navy"], 24, 3)
        w.border(p, 'left', D["gold"], 48, 0)
        p.paragraph_format.left_indent = Cm(0.25); p.paragraph_format.right_indent = Cm(0.1)
        p.paragraph_format.tab_stops.add_tab_stop(Emu(int(COLW - Cm(0.45))), WD_TAB_ALIGNMENT.RIGHT)
        r = w.font(p.add_run(f' {g["code"]} '), D["font_head"], D["size_bar"] + 1.5, True, "FFFFFF"); w.shade(r._element.get_or_add_rPr(), D["gold"])
        w.font(p.add_run('  ' + g["name"]), D["font_head"], D["size_bar"], True, "FFFFFF")
        w.font(p.add_run(f'\t{RANGE[g["code"]][0]}–{RANGE[g["code"]][1]}'), D["font_body"], 9, False, D["gold_light"])
        for n in g["questions"]:
            NEW[n] = num; q = Q[n]
            p = w.para(before=7, after=1); pf = p.paragraph_format
            pf.left_indent = IND; pf.first_line_indent = -IND; pf.keep_with_next = True; pf.tab_stops.add_tab_stop(IND)
            w.font(p.add_run(f'{num}.'), D["font_head"], D["size_q"] + 0.5, True, D["navy"])
            w.font(p.add_run('\t' + q["eng"]), D["font_body"], D["size_q"], True, D["ink"])
            if q["mm"]:
                p = w.para(after=3); p.paragraph_format.left_indent = IND; p.paragraph_format.keep_with_next = True
                w.mm(p, q["mm"], D["size_mm"])
            opts = q["opts"]; longest = max((len(o) for o in opts), default=0)
            avail = (COLW - IND) / Cm(1)                      # cm available for options
            per = 4 if longest * 0.19 + 0.8 <= avail / 4 else (2 if longest * 0.19 + 0.8 <= avail / 2 else 1)
            letters = iter('ABCD'); rows = [opts[i:i + per] for i in range(0, len(opts), per)]
            for ri, row in enumerate(rows):
                p = w.para(after=1); p.paragraph_format.left_indent = IND
                if ri < len(rows) - 1: p.paragraph_format.keep_with_next = True
                step = (COLW - IND) / per
                for k in range(1, per): p.paragraph_format.tab_stops.add_tab_stop(Emu(int(IND + step * k)))
                for k, o in enumerate(row):
                    if k: p.add_run('\t')
                    w.font(p.add_run(f'({next(letters)}) '), D["font_body"], D["size_opt"], True, D["gold"])
                    w.font(p.add_run(o), D["font_body"], D["size_opt"], color=D["ink"])
            num += 1
        w.para(after=6)

    # ---- answer key ----
    w.section(1)
    w.para('Answer Key', D["font_head"], 22, True, D["navy"], after=2)
    p = w.para(after=12); w.border(p, 'bottom', D["gold"], 12, 2); p.paragraph_format.right_indent = Cm(13)
    ordered = sorted(NEW, key=NEW.get); COLS = 8; nrows = (len(ordered) + COLS - 1) // COLS
    tbl = d.add_table(rows=nrows, cols=COLS); w.tbl_borders(tbl); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, n in enumerate(ordered):
        cc = tbl.rows[i % nrows].cells[i // nrows]; w.cell_pad(cc, 60, 60, 80, 80)
        if (i % nrows) % 2: w.shade(cc._tc.get_or_add_tcPr(), D["tint"])
        pp = cc.paragraphs[0]; pp.alignment = C
        w.font(pp.add_run(f'{NEW[n]}. '), D["font_body"], 10, False, D["muted"])
        w.font(pp.add_run(ANS.get(n, '–')), D["font_head"], 11, True, D["navy"])
    w.widths(tbl, [int(CW / COLS)] * COLS)

    if SOL:
        w.para(before=20); w.para('Worked Solutions', D["font_head"], 18, True, D["navy"], after=2)
        p = w.para(after=10); w.border(p, 'bottom', D["gold"], 12, 2); p.paragraph_format.right_indent = Cm(13)
        for g in groups:
            p = w.para(before=8, after=3); p.paragraph_format.keep_with_next = True
            w.font(p.add_run(f'{g["code"]}  '), D["font_head"], 11.5, True, D["gold"]); w.font(p.add_run(g["name"]), D["font_head"], 11.5, True, D["navy"])
            tbl = d.add_table(rows=0, cols=3); w.tbl_borders(tbl)
            for n in g["questions"]:
                r = tbl.add_row().cells
                for cc in r: w.cell_pad(cc, 50, 50, 100, 100)
                w.shade(r[0]._tc.get_or_add_tcPr(), D["tint"]); w.shade(r[1]._tc.get_or_add_tcPr(), D["tint"])
                w.font(r[0].paragraphs[0].add_run(str(NEW[n])), D["font_body"], 10, True, D["navy"])
                r[1].paragraphs[0].alignment = C
                w.font(r[1].paragraphs[0].add_run(ANS.get(n, '–')), D["font_head"], 10.5, True, D["navy"])
                w.font(r[2].paragraphs[0].add_run(SOL.get(n, '')), D["font_body"], 10, color=D["ink"])
            w.widths(tbl, [Cm(1.2), Cm(1.2), CW - Cm(2.4)])

    w.para(before=26)
    p = w.para(align=C); w.border(p, 'top', D["gold"], 8, 1); p.paragraph_format.left_indent = p.paragraph_format.right_indent = Cm(5)
    w.para(D["author"], D["font_head"], 13, True, D["navy"], C, before=8)
    w.para(D["author_role"], D["font_body"], 10, False, D["gold"], C, before=2, italic=True)

    d.core_properties.author = D["author"]; d.core_properties.title = f'{title} – {D["institute"]}'
    d.save(cfg["output"])
    no_ans = [NEW[n] for n in ordered if n not in ANS]
    print(f'Saved {cfg["output"]}: {len(Q)} questions in {len(groups)} types.' + (f' No answer for: {no_ans}' if no_ans else ''))

if __name__ == '__main__':
    ap = argparse.ArgumentParser(description='GNSI premium booklet builder')
    ap.add_argument('settings'); ap.add_argument('--columns', type=int, default=2, choices=[1, 2])
    a = ap.parse_args()
    build(json.load(open(a.settings, encoding='utf-8')), a.columns)
